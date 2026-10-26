
from pathlib import Path
import re
import random
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter


# ============================================================
# CampusMail: Manual ML Labeling Queue
# 10 categories x 10 emails = 100 candidate emails
# ============================================================

ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "outputs"

SOURCE_FILES = [
    OUTPUT_DIR / "auto_categorized_emails.csv",
    OUTPUT_DIR / "college_emails_clean.csv",
]

VALIDATION_FILE = OUTPUT_DIR / "validation_sample_50.csv"

OUTPUT_CSV = OUTPUT_DIR / "ml_labeling_queue.csv"
OUTPUT_XLSX = OUTPUT_DIR / "ml_labeling_queue.xlsx"

EMAILS_PER_CATEGORY = 10
RANDOM_SEED = 42

CATEGORIES = [
    "hackathons",
    "exams",
    "hostel_notices",
    "seminars",
    "sports",
    "scholarships",
    "assignments",
    "events",
    "general_announcements",
    "placements",
]

CATEGORY_KEYWORDS = {
    "hackathons": [
        "hackathon", "coding competition", "innovation challenge",
        "codefest", "hack fest", "programming contest",
        "coding contest", "hack day",
    ],
    "exams": [
        "examination", "exam timetable", "exam schedule",
        "hall ticket", "mid semester", "mid-semester",
        "end semester", "end-semester", "date sheet",
        "exam results", "examination notice",
    ],
    "hostel_notices": [
        "hostel", "mess menu", "hostel warden",
        "room allocation", "residential hall",
        "hostel fee", "hostel accommodation",
    ],
    "seminars": [
        "seminar", "guest lecture", "webinar",
        "technical talk", "keynote", "research talk",
        "expert talk", "conference lecture",
    ],
    "sports": [
        "sports", "tournament", "cricket match",
        "football match", "team trials", "athletics",
        "inter-college match", "sports meet",
        "badminton", "basketball tournament",
    ],
    "scholarships": [
        "scholarship", "financial aid", "financial assistance",
        "stipend", "merit-cum-means", "education grant",
        "scholarship application", "fee reimbursement",
    ],
    "assignments": [
        "assignment", "coursework", "lab report",
        "homework", "project submission",
        "submit your report", "assignment deadline",
        "laboratory assignment",
    ],
    "events": [
        "festival", "cultural event", "college fest",
        "celebration", "annual day", "concert",
        "event registration", "cultural program",
        "cultural programme", "freshers",
    ],
    "general_announcements": [
        "circular", "general notice", "administrative notice",
        "office timings", "college announcement",
        "important notice", "official announcement",
        "administrative circular",
    ],
    "placements": [
        "placement drive", "campus placement",
        "campus recruitment", "recruitment drive",
        "job opportunity", "job opening",
        "internship opportunity", "pre-placement talk",
        "pre placement talk", "placement cell",
        "on-campus hiring", "off-campus hiring",
        "career fair", "technical interview",
        "company hiring", "campus hiring",
        "placement registration", "recruitment",
        "internship", "job application",
    ],
}


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------

def normalize_text(value):
    """Normalize text for matching and duplicate detection."""
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def find_column(df, possible_names, required=True):
    """Find a column without depending on capitalization."""
    lookup = {str(col).strip().lower(): col for col in df.columns}

    for name in possible_names:
        if name.lower() in lookup:
            return lookup[name.lower()]

    # Also support columns whose names contain a known term.
    for col in df.columns:
        col_name = str(col).strip().lower()
        for name in possible_names:
            if name.lower() in col_name:
                return col

    if required:
        raise ValueError(
            f"Could not find a column matching {possible_names}.\n"
            f"Available columns: {list(df.columns)}"
        )
    return None


def email_key(subject, body):
    """A key for identifying duplicate emails."""
    return normalize_text(subject) + " || " + normalize_text(body)


def keyword_pattern(keyword):
    """Match a keyword phrase without matching inside larger words."""
    return (
        r"(?<!\w)"
        + re.escape(keyword.lower()).replace(r"\ ", r"\s+")
        + r"(?!\w)"
    )


def contains_category_keyword(text, keywords):
    text = normalize_text(text)
    return any(
        re.search(keyword_pattern(word), text)
        for word in keywords
    )


# ------------------------------------------------------------
# Load source dataset
# ------------------------------------------------------------

def load_source():
    for path in SOURCE_FILES:
        if path.exists():
            print(f"Reading emails from: {path.name}")
            df = pd.read_csv(path, low_memory=False)
            if not df.empty:
                return df

    raise FileNotFoundError(
        "Could not find a usable source CSV.\n"
        "Expected one of:\n"
        f"  {SOURCE_FILES[0]}\n"
        f"  {SOURCE_FILES[1]}"
    )


def load_excluded_keys():
    """Exclude emails in the existing 50-email validation sample."""
    excluded = set()

    if not VALIDATION_FILE.exists():
        print(
            "Warning: validation_sample_50.csv not found; "
            "no prior validation emails will be excluded."
        )
        return excluded

    sample = pd.read_csv(VALIDATION_FILE, low_memory=False)

    subject_col = find_column(
        sample, ["subject", "email_subject", "title"], required=False
    )
    body_col = find_column(
        sample,
        ["body", "email_body", "message", "content", "text"],
        required=False,
    )

    if subject_col is None and body_col is None:
        print(
            "Warning: could not identify subject/body columns in "
            "validation_sample_50.csv."
        )
        return excluded

    for _, row in sample.iterrows():
        subject = row[subject_col] if subject_col else ""
        body = row[body_col] if body_col else ""
        excluded.add(email_key(subject, body))

    print(f"Excluding {len(excluded)} previously sampled emails.")
    return excluded


# ------------------------------------------------------------
# Build candidate queue
# ------------------------------------------------------------

def build_queue():
    source = load_source()

    subject_col = find_column(
        source, ["subject", "email_subject", "title"]
    )
    body_col = find_column(
        source,
        ["body", "email_body", "message", "content", "text"],
    )

    prediction_col = find_column(
        source,
        [
            "predicted_category",
            "prediction",
            "auto_category",
            "category",
            "label",
        ],
        required=False,
    )

    id_col = find_column(
        source,
        ["email_id", "id", "message_id"],
        required=False,
    )

    # Normalize fields used by the selection process.
    source = source.copy()
    source["_subject"] = source[subject_col].fillna("").astype(str)
    source["_body"] = source[body_col].fillna("").astype(str)
    source["_key"] = source.apply(
        lambda row: email_key(row["_subject"], row["_body"]),
        axis=1,
    )

    # Remove emails without useful content.
    source = source[
        (source["_subject"].str.strip() != "")
        | (source["_body"].str.strip() != "")
    ].copy()

    # Deduplicate source emails.
    source = source.drop_duplicates(subset="_key", keep="first")

    excluded_keys = load_excluded_keys()
    source = source[~source["_key"].isin(excluded_keys)].copy()

    if len(source) < len(CATEGORIES) * EMAILS_PER_CATEGORY:
        raise ValueError(
            f"Only {len(source)} unique emails remain after exclusions. "
            f"At least {len(CATEGORIES) * EMAILS_PER_CATEGORY} are "
            "needed for 100 distinct candidates."
        )

    # Stable randomness for reproducible results.
    source = source.sample(
        frac=1, random_state=RANDOM_SEED
    ).reset_index(drop=True)

    # Combine subject and body for keyword matching.
    source["_search_text"] = (
        source["_subject"] + " " + source["_body"]
    ).map(normalize_text)

    selected_rows = []
    used_keys = set()

    for category in CATEGORIES:
        keywords = CATEGORY_KEYWORDS[category]

        # Prefer emails with relevant category keywords.
        matches = source[
            source["_search_text"].apply(
                lambda text: contains_category_keyword(text, keywords)
            )
        ]

        matches = matches[~matches["_key"].isin(used_keys)]

        chosen = matches.head(EMAILS_PER_CATEGORY).copy()

        # If too few keyword matches exist, fill from the remaining pool.
        # These are only candidates; they must be manually checked.
        if len(chosen) < EMAILS_PER_CATEGORY:
            needed = EMAILS_PER_CATEGORY - len(chosen)
            already_chosen = set(chosen["_key"])
            remaining = source[
                ~source["_key"].isin(used_keys)
                & ~source["_key"].isin(already_chosen)
            ]

            fallback = remaining.head(needed)
            chosen = pd.concat([chosen, fallback], ignore_index=False)

            print(
                f"Warning: {category} had insufficient keyword matches; "
                f"added {len(fallback)} fallback candidate(s)."
            )

        chosen = chosen.head(EMAILS_PER_CATEGORY)

        for _, row in chosen.iterrows():
            key = row["_key"]
            used_keys.add(key)

            email_id = (
                row[id_col]
                if id_col is not None
                else f"EMAIL_{len(selected_rows) + 1:04d}"
            )

            previous_prediction = (
                row[prediction_col]
                if prediction_col is not None
                else ""
            )

            selected_rows.append({
                "email_id": email_id,
                "subject": row["_subject"],
                "body": row["_body"],
                "previous_prediction": previous_prediction,
                "suggested_category": category,
                "verified_category": "",
                "review_status": "Pending",
                "review_notes": "",
            })

        print(
            f"{category}: selected {len(chosen)} candidate(s) "
            f"out of {EMAILS_PER_CATEGORY}"
        )

    queue = pd.DataFrame(selected_rows)

    if len(queue) != len(CATEGORIES) * EMAILS_PER_CATEGORY:
        raise RuntimeError(
            f"Expected {len(CATEGORIES) * EMAILS_PER_CATEGORY} rows, "
            f"but created {len(queue)}. Check the source dataset."
        )

    if queue["email_id"].duplicated().any():
        # Preserve unique IDs even if the original dataset has repeated IDs.
        queue["email_id"] = [
            f"EMAIL_{i:04d}" for i in range(1, len(queue) + 1)
        ]

    return queue


# ------------------------------------------------------------
# Create Excel workbook with dropdowns
# ------------------------------------------------------------

def create_excel(queue):
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        queue.to_excel(
            writer, sheet_name="Label Emails", index=False
        )

        instructions = pd.DataFrame({
            "Instructions": [
                "CampusMail - Manual Email Verification",
                "",
                "There are 10 categories and 10 candidate emails per category.",
                "Read each email's subject and body before assigning a label.",
                "Select the correct category in verified_category.",
                "Change review_status to Verified only after checking the email.",
                "The suggested_category is a candidate-selection hint, not ground truth.",
                "If the email does not fit a category, leave it Pending and explain in review_notes.",
                "Save this workbook after labeling.",
                "Do not train models using suggested_category as the true label.",
                "Only manually verified verified_category values should be used as ground truth.",
            ]
        })
        instructions.to_excel(
            writer, sheet_name="Instructions", index=False
        )

    wb = load_workbook(OUTPUT_XLSX)
    ws = wb["Label Emails"]

    # Freeze header and enable filters.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30

    header_fill = PatternFill(
        fill_type="solid", fgColor="1F4E78"
    )

    for cell in ws[1]:
        cell.font = Font(
            bold=True, color="FFFFFF"
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            vertical="center", wrap_text=True
        )

    # Find columns by header names.
    header_map = {
        cell.value: cell.column for cell in ws[1]
    }

    category_col = header_map["verified_category"]
    status_col = header_map["review_status"]

    category_letter = get_column_letter(category_col)
    status_letter = get_column_letter(status_col)

    last_row = ws.max_row

    # Category dropdown.
    category_values = ",".join(CATEGORIES)
    category_validation = DataValidation(
        type="list",
        formula1=f'"{category_values}"',
        allow_blank=True,
    )
    category_validation.error = (
        "Please select one of the 10 available categories."
    )
    category_validation.errorTitle = "Invalid category"
    category_validation.prompt = (
        "Choose the category that correctly describes this email."
    )
    category_validation.promptTitle = "Verify email category"
    ws.add_data_validation(category_validation)
    category_validation.add(
        f"{category_letter}2:{category_letter}{last_row}"
    )

    # Review status dropdown.
    status_validation = DataValidation(
        type="list",
        formula1='"Pending,Verified,Needs Review"',
        allow_blank=False,
    )
    ws.add_data_validation(status_validation)
    status_validation.add(
        f"{status_letter}2:{status_letter}{last_row}"
    )

    # Highlight verified rows.
    verified_fill = PatternFill(
        fill_type="solid", fgColor="D9EAD3"
    )
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(ws.max_column)}{last_row}",
        FormulaRule(
            formula=[f'${status_letter}2="Verified"'],
            fill=verified_fill,
        ),
    )

    # Make long email bodies readable.
    widths = {
        "email_id": 15,
        "subject": 40,
        "body": 65,
        "previous_prediction": 24,
        "suggested_category": 25,
        "verified_category": 25,
        "review_status": 18,
        "review_notes": 35,
    }

    for name, col_index in header_map.items():
        ws.column_dimensions[
            get_column_letter(col_index)
        ].width = widths.get(name, 20)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top", wrap_text=True
            )

    # Instructions sheet formatting.
    inst = wb["Instructions"]
    inst.column_dimensions["A"].width = 110
    inst.freeze_panes = "A2"

    for cell in inst["A"]:
        cell.alignment = Alignment(
            vertical="top", wrap_text=True
        )

    inst["A1"].font = Font(bold=True, size=14, color="1F4E78")

    wb.save(OUTPUT_XLSX)


# ------------------------------------------------------------
# Main
# ------------------------------------------------------------

def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Avoid confusing Windows file-lock errors.
    for path in [OUTPUT_CSV, OUTPUT_XLSX]:
        if path.exists():
            try:
                with open(path, "a", encoding="utf-8"):
                    pass
            except PermissionError:
                raise PermissionError(
                    f"Cannot write to {path}.\n"
                    "Close the file in Excel or another application, "
                    "then run this script again."
                )

    queue = build_queue()

    # Write the CSV first.
    queue.to_csv(
        OUTPUT_CSV,
        index=False,
        encoding="utf-8-sig",
    )

    # Create Excel workbook.
    create_excel(queue)

    print("\n" + "=" * 55)
    print("CampusMail labeling queue created successfully!")
    print("=" * 55)
    print(f"Total candidate emails: {len(queue)}")
    print(f"Categories: {len(CATEGORIES)}")
    print(f"Candidates per category: {EMAILS_PER_CATEGORY}")
    print(f"\nCSV:   {OUTPUT_CSV}")
    print(f"Excel: {OUTPUT_XLSX}")
    print("\nCandidate count by suggested category:")
    print(queue["suggested_category"].value_counts().to_string())
    print("\nNext: open the Excel file and manually verify every email.")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}")
        sys.exit(1)
